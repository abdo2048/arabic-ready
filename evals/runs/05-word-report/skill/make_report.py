from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

THEME_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme")
FONT = "Arial"

def make_rtl_document(font=FONT, size_pt=12, lang="ar-EG"):
    doc = Document()
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    rpr = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    fonts = rpr.find(qn("w:rFonts"))
    for a in THEME_ATTRS:
        fonts.attrib.pop(qn(a), None)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(a), font)
    for tag in ("w:sz", "w:szCs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); rpr.append(el)
        el.set(qn("w:val"), str(size_pt * 2))
    lang_el = rpr.find(qn("w:lang"))
    if lang_el is None:
        lang_el = OxmlElement("w:lang"); rpr.append(lang_el)
    lang_el.set(qn("w:bidi"), lang)
    rpr.append(OxmlElement("w:rtl"))
    ppr_default = defaults.find(qn("w:pPrDefault"))
    if ppr_default is None:
        ppr_default = OxmlElement("w:pPrDefault"); defaults.append(ppr_default)
    ppr = ppr_default.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr"); ppr_default.append(ppr)
    ppr.insert(0, OxmlElement("w:bidi"))
    for st in doc.styles:
        st_rpr = st.element.find(qn("w:rPr"))
        if st_rpr is not None and st_rpr.find(qn("w:rFonts")) is not None:
            f = st_rpr.find(qn("w:rFonts"))
            for a in THEME_ATTRS:
                f.attrib.pop(qn(a), None)
            for a in ("w:ascii", "w:hAnsi", "w:cs"):
                f.set(qn(a), font)
    return doc

def bold(run):
    run.bold = True
    run._r.get_or_add_rPr().append(OxmlElement("w:bCs"))

def size(run, pt):
    rpr = run._r.get_or_add_rPr()
    for tag in ("w:sz", "w:szCs"):
        el = OxmlElement(tag); el.set(qn("w:val"), str(int(pt * 2))); rpr.append(el)

def rtl_table(table):
    table._tbl.tblPr.append(OxmlElement("w:bidiVisual"))

def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)

def line_spacing(p, mult=1.5, after=6):
    pf = p.paragraph_format
    pf.line_spacing = mult
    pf.space_after = Pt(after)

AR = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")
def d(s): return str(s).translate(AR)

ACCENT = RGBColor(0x1F, 0x4E, 0x79)

doc = make_rtl_document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
for m in ("top_margin", "bottom_margin"): setattr(sec, m, Cm(2))
for m in ("left_margin", "right_margin"): setattr(sec, m, Cm(2.2))

# Title
t = doc.add_heading(level=0)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("تقرير الموارد البشرية للربع الثالث " + d(2026)); size(r, 22); bold(r)
r.font.color.rgb = ACCENT
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("شركة النيل للصناعات الغذائية – القاهرة"); size(r, 13)
r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
line_spacing(sub, 1.3, 12)

# Paragraphs
h = doc.add_heading("ملخص الربع", level=2)
p1 = doc.add_paragraph(
    "شهد الربع الثالث من عام " + d(2026) + " استقرارًا نسبيًا في القوى العاملة، إذ بلغ إجمالي عدد الموظفين "
    + d(412) + " موظفًا. وسجّل معدل دوران الموظفين " + d(12) + "٪، وهو أعلى قليلًا من المستهدف، "
    "وتركّزت معظم حالات الاستقالة في إدارة المبيعات بسبب المنافسة على الكفاءات في السوق المحلية."
)
line_spacing(p1, 1.6, 8)
p2 = doc.add_paragraph(
    "وتواصل الإدارة تنفيذ خطة التوظيف لعام " + d(2026) + "، التي تستهدف تعيين " + d(45)
    + " موظفًا جديدًا خلال الربع الأخير مع التركيز على الوظائف التقنية والتشغيلية، إلى جانب برامج استبقاء تشمل مراجعة الرواتب "
    "ومسارات واضحة للترقي الوظيفي، بهدف خفض معدل الدوران إلى أقل من " + d(9) + "٪ في عام " + d(2027) + "."
)
line_spacing(p2, 1.6, 8)

# Highlights
doc.add_heading("أبرز الإنجازات", level=2)
for item in [
    "إتمام تدريب " + d(180) + " موظفًا على برامج السلامة المهنية والتحول الرقمي.",
    "خفض متوسط مدة التوظيف من " + d(42) + " يومًا إلى " + d(30) + " يومًا.",
    "ارتفاع مؤشر رضا الموظفين إلى " + d(78) + "٪ مقارنةً بـ" + d(71) + "٪ في الربع السابق.",
]:
    bp = doc.add_paragraph(item, style="List Bullet")
    line_spacing(bp, 1.5, 4)

# Table
doc.add_heading("عدد الموظفين حسب الإدارة", level=2)
rows = [
    ("الإدارة", "عدد الموظفين", "التعيينات الجديدة", "معدل الدوران"),
    ("الإنتاج والتشغيل", 168, 9, None),
    ("المبيعات والتسويق", 112, 11, None),
    ("الشؤون المالية والإدارية", 74, 3, None),
    ("تقنية المعلومات", 58, 6, None),
]
turnover = [None, 10, 18, 8, 11]
total = ("الإجمالي", 412, 29, 12)

tbl = doc.add_table(rows=len(rows) + 1, cols=4)
tbl.style = "Table Grid"
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
rtl_table(tbl)
tbl.autofit = False
widths = [Cm(5.6), Cm(3.6), Cm(3.6), Cm(3.6)]
for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths):
    gc.set(qn("w:w"), str(int(w.twips)))

def fill(cell, text, is_bold=False, center=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
    if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text); size(r, 11)
    if is_bold: bold(r)
    if color: r.font.color.rgb = color

for i, row in enumerate(rows):
    if i == 0:
        vals = list(row)
    else:
        vals = [row[0], d(row[1]), d(row[2]), d(turnover[i]) + "٪"]
    for j, v in enumerate(vals):
        c = tbl.cell(i, j); c.width = widths[j]
        fill(c, v, is_bold=(i == 0), center=(j > 0),
             color=RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else None)
        if i == 0: shade(c, "1F4E79")
        elif i % 2 == 0: shade(c, "EAF1F8")
last = len(rows)
for j, v in enumerate([total[0], d(total[1]), d(total[2]), d(total[3]) + "٪"]):
    c = tbl.cell(last, j); c.width = widths[j]
    fill(c, v, is_bold=True, center=(j > 0)); shade(c, "D6E3F0")

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(8)
r = note.add_run("إعداد: إدارة الموارد البشرية – أكتوبر " + d(2026)); size(r, 10)
r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

doc.save("report.docx")
