from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

GREEN = RGBColor(0x0B, 0x4F, 0x3C)
GOLD = RGBColor(0xB8, 0x8A, 0x2E)
INK = RGBColor(0x1F, 0x2A, 0x2E)
MUTED = RGBColor(0x5B, 0x6B, 0x66)
TINT = RGBColor(0xEE, 0xF4, 0xF1)
FONT = "Arial"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

def rtl_para(p, align=PP_ALIGN.RIGHT):
    p.alignment = align
    p._p.get_or_add_pPr().set("rtl", "1")

def style_run(r, size, color, bold=False):
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT
    rPr = r._r.get_or_add_rPr()
    rPr.set("lang", "ar-AE"); rPr.set("altLang", "en-US")
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)

def textbox(x, y, w, h, name, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h); tb.name = name
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf

M = Inches(0.6)
W = prs.slide_width

# Title (right aligned, RTL)
_, tf = textbox(M, Inches(0.55), W - 2 * M, Inches(0.9), "Title")
p = tf.paragraphs[0]; rtl_para(p)
style_run(p.add_run(), 40, GREEN, bold=True); p.runs[0].text = "نتائج الربع الثالث"
p2 = tf.add_paragraph(); rtl_para(p2)
r = p2.add_run(); r.text = "اجتماع مجلس الإدارة"; style_run(r, 16, MUTED)

# Three stat cards, ordered right-to-left
stats = [
    ("25%", "نمو الإيرادات", "نمو الإيرادات بنسبة 25%"),
    ("1,200", "عميل جديد", "انضمام 1,200 عميل جديد"),
    ("62", "مؤشر صافي الترويج (NPS)", "تحقيق 62 في مؤشر صافي الترويج (NPS)"),
]
gap = Inches(0.4)
card_w = int((W - 2 * M - 2 * gap) / 3)
card_h = Inches(2.3)
top = Inches(2.0)
for i, (num, label, _) in enumerate(stats):
    x = W - M - card_w - i * (card_w + gap)  # first card on the right
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, top, card_w, card_h)
    card.name = f"Stat card {i+1}"
    card.adjustments[0] = 0.08
    card.fill.solid(); card.fill.fore_color.rgb = TINT; card.line.fill.background()
    card.shadow.inherit = False
    _, ctf = textbox(x + Inches(0.3), top + Inches(0.25), card_w - Inches(0.6), card_h - Inches(0.5),
                     f"Stat text {i+1}", MSO_ANCHOR.MIDDLE)
    p = ctf.paragraphs[0]; rtl_para(p, PP_ALIGN.CENTER)
    r = p.add_run(); r.text = num; style_run(r, 60, GREEN, bold=True)
    p = ctf.add_paragraph(); rtl_para(p, PP_ALIGN.CENTER)
    r = p.add_run(); r.text = label; style_run(r, 18, INK)

# Bullet list (real bullets, RTL)
_, btf = textbox(M, Inches(4.75), W - 2 * M, Inches(2.1), "Key points")
for i, (_, _, sentence) in enumerate(stats):
    p = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
    rtl_para(p)
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.4))); pPr.set("indent", str(-Inches(0.3)))
    sa = etree.SubElement(pPr, qn("a:spcAft")); etree.SubElement(sa, qn("a:spcPts")).set("val", "1000")
    bc = etree.SubElement(pPr, qn("a:buClr")); etree.SubElement(bc, qn("a:srgbClr")).set("val", "B88A2E")
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
    r = p.add_run(); r.text = sentence; style_run(r, 20, INK)

prs.save("slide.pptx")
print("saved slide.pptx")
