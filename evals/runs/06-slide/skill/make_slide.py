from pptx import Presentation
from pptx.util import Pt, Emu, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

FONT, LANG = "Arial", "ar-AE"
NAVY, GOLD, INK = RGBColor(0x0B, 0x2A, 0x4A), RGBColor(0xB8, 0x8A, 0x2E), RGBColor(0x22, 0x2B, 0x36)

def rtl_paragraph(p, bullet=False):
    pPr = p._p.get_or_add_pPr()
    pPr.set("rtl", "1")
    pPr.set("algn", "r")
    if bullet:
        pPr.set("marL", str(Emu(Pt(34))))   # start-side indent (right side in RTL)
        pPr.set("indent", str(-Emu(Pt(34))))
        for tag in ("a:buClr", "a:buFont", "a:buChar"):
            old = pPr.find(qn(tag))
            if old is not None: pPr.remove(old)
        clr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr")).set("val", "B88A2E")
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", FONT)
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
    else:
        etree.SubElement(pPr, qn("a:buNone"))
    for r in p.runs:
        rPr = r._r.get_or_add_rPr()
        rPr.set("lang", LANG)
        for tag in ("a:latin", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rPr, qn(tag))
            el.set("typeface", FONT)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])
W = prs.slide_width

# accent bar on the start (right) side of the title
bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, W - Inches(0.9), Inches(0.9), Inches(0.12), Inches(1.0))
bar.fill.solid(); bar.fill.fore_color.rgb = GOLD; bar.line.fill.background()

title = slide.shapes.add_textbox(Inches(0.9), Inches(0.75), W - Inches(2.0), Inches(1.3)).text_frame
title.word_wrap = True; title.vertical_anchor = MSO_ANCHOR.MIDDLE
p = title.paragraphs[0]; r = p.add_run(); r.text = "نتائج الربع الثالث"
r.font.size = Pt(44); r.font.bold = True; r.font.color.rgb = NAVY
rtl_paragraph(p)

rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(2.25), W - Inches(1.8), Pt(1.5))
rule.fill.solid(); rule.fill.fore_color.rgb = RGBColor(0xD9, 0xDE, 0xE5); rule.line.fill.background()

body = slide.shapes.add_textbox(Inches(0.9), Inches(2.7), W - Inches(1.8), Inches(3.8)).text_frame
body.word_wrap = True
bullets = [
    "نمو الإيرادات بنسبة 25%",
    "استقطاب 1,200 عميل جديد",
    "وصول مؤشر صافي نقاط الترويج (NPS) إلى 62",
]
for i, text in enumerate(bullets):
    p = body.paragraphs[0] if i == 0 else body.add_paragraph()
    r = p.add_run(); r.text = text
    r.font.size = Pt(30); r.font.color.rgb = INK
    p.line_spacing = 1.3
    p.space_after = Pt(18)
    rtl_paragraph(p, bullet=True)

prs.save("slide.pptx")
