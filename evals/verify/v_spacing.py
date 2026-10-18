# letter-spacing on Arabic: LibreOffice (docx w:spacing) and Pillow has none; Chromium ignores it (see web_typo.png)
import subprocess
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from common import *
d = Document()
for sp in (0, 60):
    p = d.add_paragraph(); p._p.get_or_add_pPr().insert(0, OxmlElement("w:bidi"))
    r = p.add_run(f"مرحبا بكم في متجرنا (spacing={sp})")
    rPr = r._r.get_or_add_rPr()
    f = OxmlElement("w:rFonts"); f.set(qn("w:cs"), "Noto Naskh Arabic"); rPr.append(f)
    s = OxmlElement("w:spacing"); s.set(qn("w:val"), str(sp)); rPr.append(s)
    for t in ("w:sz", "w:szCs"):
        e = OxmlElement(t); e.set(qn("w:val"), "36"); rPr.append(e)
    rPr.append(OxmlElement("w:rtl"))
d.save(f"{OUT}/docx_letterspacing.docx")
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", OUT, f"{OUT}/docx_letterspacing.docx"], capture_output=True)
import pymupdf
pg = pymupdf.open(f"{OUT}/docx_letterspacing.pdf")[0]
pg.get_pixmap(matrix=pymupdf.Matrix(2, 2), clip=pymupdf.Rect(0, 50, pg.rect.width, 160)).save(f"{OUT}/docx_letterspacing.png")
