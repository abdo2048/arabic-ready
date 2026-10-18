# Word kashida justification (w:jc lowKashida/highKashida) rendered by LibreOffice
import subprocess
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from common import *
LONG = "هذا نص عربي طويل لاختبار ضبط الأسطر بالكشيدة في مستندات وورد عندما يكون العمود ضيقا والنص طويلا بما يكفي لعدة أسطر."
d = Document()
sec = d.sections[0]
from docx.shared import Cm
sec.page_width = Cm(10); sec.left_margin = sec.right_margin = Cm(1)
for jc in ("both", "highKashida"):
    p = d.add_paragraph(); pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:bidi"))
    j = OxmlElement("w:jc"); j.set(qn("w:val"), jc); pPr.append(j)
    r = p.add_run(f"[{jc}] " + LONG); rPr = r._r.get_or_add_rPr()
    f = OxmlElement("w:rFonts"); f.set(qn("w:cs"), "Noto Naskh Arabic"); rPr.append(f)
    for t in ("w:sz", "w:szCs"):
        e = OxmlElement(t); e.set(qn("w:val"), "28"); rPr.append(e)
    rPr.append(OxmlElement("w:rtl"))
d.save(f"{OUT}/docx_kashida.docx")
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", OUT, f"{OUT}/docx_kashida.docx"], capture_output=True)
import pymupdf
pg = pymupdf.open(f"{OUT}/docx_kashida.pdf")[0]
pg.get_pixmap(matrix=pymupdf.Matrix(2, 2), clip=pymupdf.Rect(0, 20, pg.rect.width, 330)).save(f"{OUT}/docx_kashida.png")
